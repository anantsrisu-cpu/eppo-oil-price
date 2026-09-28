"""Offline unit tests (no network). Run:  python -m pytest -q   or   python -m unittest -v"""
from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from eppo import api, build, config, structure, validate  # noqa: E402
from eppo.http import _check_url, _looks_blocked  # noqa: E402
from eppo.storage import read_csv, upsert  # noqa: E402

FIXTURE = ROOT / "data" / "raw" / "api" / "2026" / "09" / "2026-09-25.json"


def make_structure_xlsx() -> bytes:
    """Synthetic workbook with the same layout as EPPO pt-price-st-*.xlsx."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Oil Price Structure"
    ws["B2"] = "PRICE STRUCTURE OF PETROLEUM PRODUCT IN BANGKOK"
    ws["B4"] = 46290  # Excel serial = 2026-09-25
    hdr = ["UNIT: BAHT/LITRE", "EX-REFIN. ", "DISCOUNT", "EXCISE TAX ", "M. TAX ", "OIL FUND ",
           "CONSV. FUND ", "WHOLESALE (WS)", "VAT (WS)", "WS&VAT ", "MARKETING MARGIN ", "VAT (MM)", "RETAIL "]
    for j, h in enumerate(hdr):
        ws.cell(row=6, column=2 + j, value=h)
    ws.append([None, "GASOHOL95 E10", 30.707086, None, 6.75, 0.675, -4.2, 0.05, 33.9821, 2.378747,
               36.3608, 3.3450467, 0.2341533, 39.94])
    ws.append([None, "H-DIESEL ", 38.0129614, -4, 6.92, 0.692, -4.95, 0.05, 36.725, 2.57075,
               39.2958, 2.0039252, 0.1402748, 41.44])
    ws.append([None, "H-DIESEL 20", 38.4467, -4, 5.953, 0.5953, -8.99, 0.05, 32.055, 2.24385,
               34.2989, 2.0010280, 0.1400720, 36.44])
    ws.append([])
    ws.append([None, "Exchange Rate", "=", 33.5832, "BAHT/USD"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestApi(unittest.TestCase):
    def setUp(self):
        self.raw = json.loads(FIXTURE.read_text(encoding="utf-8"))

    def test_parse_all_brands(self):
        rows, warns = api.parse(self.raw, "2026-09-25", "t")
        self.assertEqual(len(rows), 64)
        self.assertEqual({r["brand_code"] for r in rows}, set(config.BRANDS))
        ptt95 = next(r for r in rows if r["brand_code"] == "ptt" and r["product_code"] == "gh95")
        self.assertEqual(ptt95["price"], 39.94)
        self.assertEqual(ptt95["effective_date"], "2026-09-24")
        self.assertTrue(any("esso" in w and "stale" in w for w in warns))

    def test_dash_is_missing(self):
        rows, _ = api.parse(self.raw, "2026-09-25", "t")
        self.assertFalse(any(r["brand_code"] == "bcp" and r["product_code"] == "gl95" for r in rows))

    def test_bad_status(self):
        with self.assertRaises(ValueError):
            api.parse({"status": "error", "data": {}}, "2026-09-25", "t")

    def test_new_brand_warning(self):
        raw = json.loads(json.dumps(self.raw))
        raw["data"]["newco"] = {"oil_newco_gh95": "40.00", "oil_newco_date": "2026-09-24"}
        _, warns = api.parse(raw, "2026-09-25", "t")
        self.assertTrue(any("NEW brand" in w for w in warns))


class TestValidate(unittest.TestCase):
    def test_ok(self):
        rows, _ = api.parse(json.loads(FIXTURE.read_text(encoding="utf-8")), "2026-09-25", "t")
        validate.check_brand_rows(rows)

    def test_out_of_range(self):
        rows, _ = api.parse(json.loads(FIXTURE.read_text(encoding="utf-8")), "2026-09-25", "t")
        rows[0]["price"] = 999
        with self.assertRaises(validate.ValidationError):
            validate.check_brand_rows(rows)

    def test_too_few(self):
        with self.assertRaises(validate.ValidationError):
            validate.check_brand_rows([])

    def test_jump_warning(self):
        rows, _ = api.parse(json.loads(FIXTURE.read_text(encoding="utf-8")), "2026-09-25", "t")
        prev = {("ptt", "gh95"): 30.0}
        self.assertTrue(validate.check_brand_rows(rows, prev))


class TestStructure(unittest.TestCase):
    def test_parse_workbook(self):
        rows = structure.parse_workbook(make_structure_xlsx(), "2026-09-25", "https://x/pt-price-st-2026-9-25.xlsx")
        by = {r["product_code"]: r for r in rows}
        self.assertEqual(set(by), {"gh95", "ds", "dsb20"})       # 'H-DIESEL 20' typo mapped
        self.assertEqual(by["gh95"]["retail"], 39.94)
        self.assertEqual(by["ds"]["refinery_discount"], -4)
        self.assertEqual(by["gh95"]["fx_thb_usd"], 33.5832)
        self.assertFalse(validate.check_structure_rows(rows))  # components add up to retail

    def test_date_from_url(self):
        self.assertEqual(structure.date_from_url(".../pt-price-st-2026-9-5.xlsx"), "2026-09-05")
        self.assertIsNone(structure.date_from_url(".../pt-price-st-2019-076-31.xls"))

    def test_rejects_html(self):
        with self.assertRaises(ValueError):
            structure.parse_workbook(b"<html>blocked</html>", "2026-09-25")


class TestHttpSecurity(unittest.TestCase):
    def test_allow_list(self):
        _check_url("https://www.eppo.go.th/wp-json/oil-api/v1/oil-prices")
        for bad in ("http://www.eppo.go.th/", "https://evil.example.com/", "https://www.eppo.go.th.evil.com/",
                    "file:///etc/passwd"):
            with self.assertRaises(ValueError):
                _check_url(bad)

    def test_cloudflare_detect(self):
        self.assertTrue(_looks_blocked(b"<title>Attention Required! | Cloudflare</title>", "text/html"))
        self.assertFalse(_looks_blocked(b'{"status":"success"}', "application/json"))


class TestStorageAndBuild(unittest.TestCase):
    def test_upsert_idempotent(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "x.csv"
            rows = [{"k": "1", "v": "a"}, {"k": "2", "v": "b"}]
            self.assertEqual(upsert(p, rows, ["k", "v"], ("k",)), (2, 0))
            self.assertEqual(upsert(p, rows, ["k", "v"], ("k",)), (0, 0))
            self.assertEqual(upsert(p, [{"k": "2", "v": "c"}], ["k", "v"], ("k",)), (0, 1))
            self.assertEqual([r["v"] for r in read_csv(p)], ["a", "c"])

    def test_daily_ffill_and_priority(self):
        structure_rows = [
            {"date": "2026-09-18", "product_code": "gh95", "retail": "39.00", "source_file": "f"},
            {"date": "2026-09-22", "product_code": "gh95", "retail": "39.50", "source_file": "f"},
        ]
        brand_rows = [{"snapshot_date": "2026-09-25", "brand_code": "ptt", "product_code": "gh95",
                       "price": "39.94", "effective_date": "2026-09-24", "effective_time": "05:00"}]
        daily = build.build_daily(brand_rows, structure_rows, "2026-09-25")
        got = {r["date"]: (r["price"], r["source"]) for r in daily}
        self.assertEqual(got["2026-09-19"], (39.0, "carried_forward"))      # weekend ffill
        self.assertEqual(got["2026-09-23"], (39.5, "carried_forward"))
        self.assertEqual(got["2026-09-24"], (39.94, "eppo_api_effective"))  # from API effective date
        self.assertEqual(got["2026-09-25"], (39.94, "eppo_api"))

    def test_stale_brand_excluded(self):
        brand_rows = [{"snapshot_date": "2026-09-25", "brand_code": "esso", "product_code": "gh95",
                       "price": "32.65", "effective_date": "2025-09-24", "effective_time": "05:00"}]
        self.assertEqual(build.stale_brands(brand_rows), {"esso": "2025-09-24"})
        self.assertEqual(build.build_daily(brand_rows, [], "2026-09-25"), [])

    def test_manual_history_change_log(self):
        with tempfile.TemporaryDirectory() as d:
            old = config.DATA_DIR
            try:
                config.DATA_DIR = Path(d)
                (Path(d) / "manual").mkdir()
                (Path(d) / "manual" / "bcp_history.csv").write_text(
                    "date,pds,ds,dsb20,gs95p,e85,e20,gh91,gh95\n01/01/2568,,32.94,,,,,,36.25\n2025-03-01,,31.94,,,,,,35.25\nbad,1,2\n",
                    encoding="utf-8")
                rows, warns = build.read_manual_history()
            finally:
                config.DATA_DIR = old
        self.assertEqual(len(rows), 4)
        self.assertTrue(any("อ่านวันที่ไม่ได้" in w for w in warns))
        daily = build.build_daily([], [], "2025-03-02", rows)
        got = {(r["date"], r["product_code"]): r["price"] for r in daily}
        self.assertEqual(got[("2025-02-20", "ds")], 32.94)   # carried > 10 days (change log)
        self.assertEqual(got[("2025-03-02", "gh95")], 35.25)

    def test_formula_injection_guard(self):
        self.assertEqual(build._safe_cell("=HYPERLINK(1)"), "'=HYPERLINK(1)")
        self.assertEqual(build._safe_cell(-4.2), -4.2)


if __name__ == "__main__":
    unittest.main()
