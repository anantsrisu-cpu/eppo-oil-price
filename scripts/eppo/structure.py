"""Daily oil price-structure Excel files (โครงสร้างราคาขายปลีกน้ำมัน).

EPPO publishes one file per business day, e.g.
  https://www.eppo.go.th/wp-content/uploads/2026/09/pt-price-st-2026-9-25.xlsx
Sheet 'Oil Price Structure' has a header row (EX-REFIN. ... RETAIL) and one row per product.
The RETAIL column = retail price of the market leader (PTT) in Bangkok.
Files exist from 2002 to today (old years are .xls, newer .xlsx; EPPO re-uploaded them in 2026).
We use them for PTT history from 2018-01-01; the per-brand archive (archive.py) covers
all 10 brands but ends on 10 Jul 2018.
"""
from __future__ import annotations

import hashlib
import io
import logging
import re
import time
from datetime import date, datetime, timedelta

from . import config
from .http import get_bytes, get_json

log = logging.getLogger(__name__)
_NAME = re.compile(r"pt-price-st-?(\d{4})[-_](\d{1,2})[-_](\d{1,2})\.xlsx?$", re.I)


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().upper()


def date_from_url(url: str) -> str | None:
    m = _NAME.search(url)
    if not m:
        return None
    try:
        return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
    except ValueError:
        return None


def list_files(since: str, max_pages: int = 80, per_page: int = 100) -> dict[str, str]:
    """Return {YYYY-MM-DD: url} for structure files dated >= since (newest first paging)."""
    found: dict[str, str] = {}
    for page in range(1, max_pages + 1):
        items, _ = get_json(config.MEDIA_API, {
            "search": config.STRUCTURE_SEARCH, "per_page": per_page, "page": page,
            "orderby": "date", "order": "desc", "_fields": "source_url,date",
        })
        if not isinstance(items, list) or not items:
            break
        for it in items:
            url = str(it.get("source_url", ""))
            d = date_from_url(url)
            if d and d >= since and url.lower().endswith((".xlsx", ".xls")):
                # keep the first (= most recently uploaded) file for each date
                found.setdefault(d, url)
        # NOTE: media is ordered by *upload* date, and EPPO bulk-uploaded old years in 2026,
        # so we cannot stop early on old file dates - we scan up to max_pages.
        if len(items) < per_page:
            break
        time.sleep(config.POLITE_DELAY)
    return dict(sorted(found.items()))


def _read_grids(content: bytes) -> list[list[list]]:
    """Return every sheet as a list of rows. Supports .xlsx (openpyxl) and old .xls (xlrd)."""
    if content[:2] == b"PK":
        from openpyxl import load_workbook  # imported lazily so the API job works without it
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        grids = []
        for ws in wb.worksheets:
            ws.reset_dimensions()  # some Excel files declare a wrong used-range
            grids.append([list(r) for r in ws.iter_rows(values_only=True)])
        wb.close()
        return grids
    if content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":  # OLE2 = Excel 97-2003 .xls
        import xlrd  # imported lazily
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            return [[sh.row_values(i) for i in range(sh.nrows)] for sh in book.sheets()]
        finally:
            book.release_resources()
    raise ValueError("not an Excel file (.xlsx / .xls)")


def _is_retail(c) -> bool:
    return _norm(c).rstrip("*").strip() == "RETAIL"


def _header_key(top, below, seen: set[str]) -> str | None:
    """Map a header cell (+ the cell under it, old files split names over 2 rows) to a field."""
    k = config.STRUCTURE_COLUMN_MAP.get(_norm(top).rstrip("*").strip())
    if k:
        return k
    t = (_norm(top) + " " + _norm(below)).strip()
    if not _norm(top):
        return None
    if t.startswith("EX-REFIN"):
        return "ex_refinery"
    if "DISCOUNT" in t:
        return "refinery_discount"
    if t.startswith("M. TAX") or t.startswith("M.TAX"):
        return "municipal_tax"
    if t.startswith("TAX") or "EXCISE" in t:
        return "excise_tax"
    if t.startswith("OIL"):
        return "oil_fund"
    if t.startswith("CONSV"):
        return "conservation_fund"
    if t.startswith("WS") and "VAT" in t:
        return "wholesale_incl_vat"
    if t.startswith("WHOLESALE"):
        return "wholesale"
    if t.startswith("MARKETING"):
        return "marketing_margin"
    if t.startswith("VAT"):
        return "vat_marketing_margin" if "marketing_margin" in seen else "vat_wholesale"
    return None


def parse_workbook(content: bytes, file_date: str, source_url: str = "") -> list[dict]:
    """Parse one structure workbook into rows (header-driven, tolerant to column moves).

    Handles every layout seen 2018-2026: .xls and .xlsx, header split over two rows
    ('OIL' / 'FUND'), 'RETAIL*', extra report sheets before the structure sheet, and the
    meaning of 'H-DIESEL' (= base diesel: B7, or B10 while a separate 'H-DIESEL B7' row exists).
    """
    return parse_grids(_read_grids(content), file_date, source_url)


def parse_grids(grids: list[list[list]], file_date: str, source_url: str = "") -> list[dict]:
    grid = header_idx = None
    for g in grids:
        idx = next((i for i, r in enumerate(g) if any(_is_retail(c) for c in r)), None)
        if idx is not None:
            grid, header_idx = g, idx
            break
    if grid is None:
        raise ValueError("header row with 'RETAIL' not found (file has no retail price)")
    head = grid[header_idx]
    below = grid[header_idx + 1] if header_idx + 1 < len(grid) else []
    colmap: dict[int, str] = {}
    label_col = None
    for j, c in enumerate(head):
        key = _header_key(c, below[j] if j < len(below) else "", set(colmap.values()))
        if key and key not in colmap.values():
            colmap[j] = key
        elif "UNIT" in _norm(c) and label_col is None:
            label_col = j
    if "retail" not in colmap.values():
        raise ValueError("RETAIL column missing")
    label_col = label_col if label_col is not None else 0

    # sheet date (Excel serial or datetime above the header) - used as a cross-check only
    sheet_date = None
    for r in grid[:header_idx]:
        for c in r:
            if isinstance(c, datetime):
                sheet_date = c.date().isoformat()
            elif isinstance(c, (int, float)) and 40000 < c < 60000:
                sheet_date = (date(1899, 12, 30) + timedelta(days=int(c))).isoformat()
    if sheet_date and sheet_date != file_date:
        log.warning("%s: sheet date %s differs from file name date", file_date, sheet_date)

    labels = []
    for r in grid[header_idx + 1:]:
        label = _norm(r[label_col] if label_col < len(r) else "")
        if label.startswith("EXCHANGE"):
            break
        labels.append((label, r))
    has_b7_row = any(lbl == "H-DIESEL B7" for lbl, _ in labels)

    rows, fx = [], None
    for r in grid[header_idx + 1:]:
        if _norm(r[label_col] if label_col < len(r) else "").startswith("EXCHANGE"):
            fx = next((c for c in r if isinstance(c, (int, float)) and c), None)
            break
    for label, r in labels:
        if not label or not any(isinstance(r[j], (int, float)) for j in colmap if j < len(r)):
            continue  # blank line or the 2nd header line
        if label == "H-DIESEL" and has_b7_row:
            code = "dsb10"   # Oct 2020 - Sep 2023: base diesel was B10, B7 listed separately
        else:
            code = config.STRUCTURE_PRODUCT_MAP.get(label)
        if code is None:
            log.warning("%s: unknown structure row %r (add to STRUCTURE_PRODUCT_MAP)", file_date, label)
            code = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")
        rec = {"date": file_date, "product_code": code, "product_label": label}
        for j, key in colmap.items():
            v = r[j] if j < len(r) else None
            rec[key] = round(float(v), 4) if isinstance(v, (int, float)) and not isinstance(v, bool) else ""
        rows.append(rec)
    for rec in rows:
        rec["fx_thb_usd"] = round(float(fx), 4) if isinstance(fx, (int, float)) else ""
        rec["source_file"] = source_url.rsplit("/", 1)[-1]
    if not rows:
        raise ValueError("no product rows found")
    return rows


def fetch_and_parse(file_date: str, url: str) -> tuple[list[dict], str]:
    content, _ = get_bytes(url, accept="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/vnd.ms-excel,*/*")
    return parse_workbook(content, file_date, url), hashlib.sha256(content).hexdigest()
