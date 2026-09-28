"""Build everything downstream of the CSV data:

  1. daily price table (calendar days x brand x product) with source flags
  2. SQLite database  -> _site/downloads/oil_prices.sqlite   (for SQL queries)
  3. pivot CSVs       -> reports/pivots/*.csv                  (committed, diff-able)
  4. Excel workbook   -> _site/downloads/oil_price_report.xlsx (pivot sheets + data table)
  5. dashboard JSON   -> _site/data/dashboard.json             (read by the web dashboard)
  6. daily report     -> reports/latest.md, reports/daily/YYYY-MM-DD.md
"""
from __future__ import annotations

import csv
import io
import json
import logging
import os
import shutil
import sqlite3
from collections import defaultdict
from datetime import date, timedelta

from . import config
from .storage import read_csv, write_csv

log = logging.getLogger(__name__)

MAX_FFILL_DAYS = 10        # carry a price forward over weekends / holidays / missed runs
MAX_EFFECTIVE_BACKFILL = 31
SOURCE_RANK = {"eppo_api": 6, "eppo_api_effective": 5, "manual": 4, "eppo_structure": 3, "seed_structure": 2,
               "eppo_archive": 1}
SRC_CODE = {"eppo_api": "a", "eppo_api_effective": "e", "eppo_structure": "s", "manual": "m",
            "seed_structure": "s", "eppo_archive": "r", "carried_forward": "f"}
# sources that list only the days a price CHANGED -> carry forward until the next change
CHANGE_LOG_SOURCES = {"manual"}


def _d(s: str) -> date:
    return date.fromisoformat(s)


def _days(a: str, b: str):
    d, end = _d(a), _d(b)
    while d <= end:
        yield d.isoformat()
        d += timedelta(days=1)


# ------------------------------------------------------------------ 1. daily table
def stale_brands(brand_rows: list[dict]) -> dict[str, str]:
    """brand -> last effective date, for brands whose latest snapshot is stale."""
    latest: dict[str, dict] = {}
    for r in brand_rows:
        if r["brand_code"] not in latest or r["snapshot_date"] > latest[r["brand_code"]]["snapshot_date"]:
            latest[r["brand_code"]] = r
    out = {}
    for b, r in latest.items():
        try:
            if (_d(r["snapshot_date"]) - _d(r["effective_date"])).days > config.STALE_AFTER_DAYS:
                out[b] = r["effective_date"]
        except ValueError:
            pass
    return out


def archive_events(rows: list[dict]) -> tuple[dict[tuple[str, str], list[tuple[str, float | None]]], str | None]:
    """EPPO per-brand archive (one snapshot per price announcement) -> change events per
    (brand, product): (date, price) when a price takes effect, (date, None) when the brand
    stops selling it. Returns (events, last announce date = end of the archive)."""
    files: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        files[r["announce_date"]].append(r)
    state: dict[tuple[str, str], tuple[str, float]] = {}
    events: dict[tuple[str, str], list[tuple[str, float | None]]] = defaultdict(list)
    for ad in sorted(files):
        present = set()
        brands_in_file = {r["brand_code"] for r in files[ad]}
        for r in files[ad]:
            k = (r["brand_code"], r["product_code"])
            price = _num(r.get("price"))
            if price is None:
                continue
            present.add(k)
            eff = r.get("effective_date") or ad
            last = state.get(k)
            if last is None or last[1] != price:
                if last is not None and eff <= last[0]:
                    eff = ad                       # inconsistent effective date -> announce date
                eff = min(eff, ad)
                events[k].append((eff, price))
                state[k] = (eff, price)
        for k in [k for k in state if k[0] in brands_in_file and k not in present]:
            events[k].append((ad, None))          # listed as 0 / missing -> not sold any more
            del state[k]
    return events, (max(files) if files else None)


def build_daily(brand_rows: list[dict], structure_rows: list[dict], end_date: str,
                manual_rows: list[dict] | None = None, archive_rows: list[dict] | None = None) -> list[dict]:
    obs: dict[tuple[str, str], dict[str, tuple[float, str]]] = defaultdict(dict)
    stops: dict[tuple[str, str], set[str]] = defaultdict(set)

    def put(b, p, d, price, src):
        cur = obs[(b, p)].get(d)
        if cur is None or SOURCE_RANK[src] > SOURCE_RANK[cur[1]]:
            obs[(b, p)][d] = (float(price), src)

    ref = config.STRUCTURE_REFERENCE_BRAND
    for r in structure_rows:
        if r["product_code"] in config.PRODUCTS and r.get("retail") not in ("", None):
            src = "seed_structure" if r.get("source_file", "").startswith("seed") else "eppo_structure"
            put(ref, r["product_code"], r["date"], r["retail"], src)

    for r in manual_rows or []:
        put(r["brand_code"], r["product_code"], r["date"], r["price"], "manual")

    arch_events, archive_end = archive_events(archive_rows or [])
    for (b, p), evs in arch_events.items():
        if b not in config.BRANDS or p not in config.PRODUCTS:
            continue
        for d, price in evs:
            if price is None:
                stops[(b, p)].add(d)
            else:
                put(b, p, d, price, "eppo_archive")

    stale = stale_brands(brand_rows)
    for r in brand_rows:
        b, p, snap = r["brand_code"], r["product_code"], r["snapshot_date"]
        if b in stale:
            # e.g. Esso: EPPO keeps returning a price that took effect a year ago (stations were
            # rebranded). Keep it in the raw snapshot table only, not in daily/pivot numbers.
            continue
        put(b, p, snap, r["price"], "eppo_api")
        # the API says since when the price is in effect -> fill those days too (bounded)
        eff = r.get("effective_date") or ""
        if b not in stale and eff and eff < snap:
            try:
                if (_d(snap) - _d(eff)).days <= MAX_EFFECTIVE_BACKFILL:
                    for d in _days(eff, (_d(snap) - timedelta(days=1)).isoformat()):
                        put(b, p, d, r["price"], "eppo_api_effective")
            except ValueError:
                pass

    out = []
    for (b, p), series in obs.items():
        start = min(series)
        last_val, last_real, last_src = None, None, None
        for d in _days(start, end_date):
            if d in series:
                price, src = series[d]
                last_val, last_real, last_src = price, d, src
            elif d in stops.get((b, p), ()):
                last_val = None
                continue
            elif last_val is not None and (
                    last_src in CHANGE_LOG_SOURCES
                    # archive = one row per price change -> valid until the next change, but never
                    # past the end of the archive (10 Jul 2018): later prices are unknown
                    or (last_src == "eppo_archive" and d <= archive_end)
                    or (last_src != "eppo_archive" and (_d(d) - _d(last_real)).days <= MAX_FFILL_DAYS)):
                price, src = last_val, "carried_forward"
            else:
                continue
            if d < config.HISTORY_START:
                continue
            out.append({"date": d, "brand_code": b, "product_code": p, "price": round(price, 2),
                        "source": src, "is_stale": 1 if b in stale and d >= stale[b] else 0})
    out.sort(key=lambda r: (r["date"], config.BRANDS.get(r["brand_code"], {}).get("order", 99),
                            config.PRODUCTS.get(r["product_code"], {}).get("order", 99)))
    return out


# ------------------------------------------------------------------ 1b. manual history import
MANUAL_WIDE_COLUMNS = {  # header aliases accepted in data/manual/<brand>_history.csv
    "pds": "pds", "ds": "ds", "b7": "ds", "dsb20": "dsb20", "b20": "dsb20", "gs95p": "gs95p", "e85": "e85",
    "e20": "e20", "gh91": "gh91", "91": "gh91", "gh95": "gh95", "95": "gh95", "gl95": "gl95", "dsb10": "dsb10",
    "gs99p": "gs99p", "gl95p": "gl95p",
}


def _parse_any_date(s: str) -> str | None:
    """Accept 2025-01-31, 31/01/2025, 31/01/2568 (Buddhist year), 31-01-68."""
    s = str(s or "").strip()
    try:
        if len(s) == 10 and s[4] == "-":
            return date.fromisoformat(s).isoformat()
        parts = s.replace("-", "/").replace(".", "/").split("/")
        dd, mm, yy = int(parts[0]), int(parts[1]), int(parts[2])
        if yy < 100:
            yy += 2500 if yy > 40 else 2000
        if yy > 2400:
            yy -= 543
        return date(yy, mm, dd).isoformat()
    except (ValueError, IndexError):
        return None


def read_manual_history() -> tuple[list[dict], list[str]]:
    """Optional price history typed/pasted by the user (e.g. copied from Bangchak's
    'ราคาน้ำมันย้อนหลัง' page). Files: data/manual/<brand_code>_history.csv
    Row = a day the price changed; columns: date + product codes (see data/manual/README.md)."""
    rows, warnings = [], []
    folder = config.DATA_DIR / "manual"
    if not folder.exists():
        return rows, warnings
    for f in sorted(folder.glob("*_history.csv")):
        brand = f.name.split("_history")[0].lower()
        if brand not in config.BRANDS:
            warnings.append(f"{f.name}: ไม่รู้จักรหัสแบรนด์ '{brand}' (ใช้ ptt, bcp, shell, caltex, irpc, pt, susco1, pure, susco2)")
            continue
        n = 0
        raw = f.read_bytes()
        for enc in ("utf-8-sig", "cp874"):   # Excel "CSV UTF-8" or Thai Windows "CSV"
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        for i, r in enumerate(csv.DictReader(io.StringIO(text)), start=2):
            cols = {str(k or "").strip().lower(): v for k, v in r.items()}
            d = _parse_any_date(cols.get("date") or cols.get("วันที่") or "")
            if not d:
                if any(str(v or "").strip() for v in cols.values()):
                    warnings.append(f"{f.name} แถว {i}: อ่านวันที่ไม่ได้ — ข้ามแถวนี้")
                continue
            for k, v in cols.items():
                code = MANUAL_WIDE_COLUMNS.get(k)
                if not code:
                    continue
                try:
                    price = float(str(v).replace(",", "").strip())
                except ValueError:
                    continue
                if not (config.PRICE_MIN <= price <= config.PRICE_MAX):
                    warnings.append(f"{f.name} แถว {i}: ราคา {price} ผิดปกติ — ข้าม")
                    continue
                rows.append({"date": d, "brand_code": brand, "product_code": code, "price": price})
                n += 1
        log.info("manual history %s: %d prices", f.name, n)
    return rows, warnings


# ------------------------------------------------------------------ 2. SQLite
SCHEMA = """
CREATE TABLE dim_brand   (brand_code TEXT PRIMARY KEY, name_th TEXT, name_en TEXT, sort_order INT,
                          is_stale INT, stale_since TEXT);
CREATE TABLE dim_product (product_code TEXT PRIMARY KEY, name_th TEXT, name_en TEXT, product_group TEXT,
                          sort_order INT);
CREATE TABLE fact_retail_daily (date TEXT, year INT, month TEXT, brand_code TEXT, product_code TEXT,
                          price REAL, source TEXT, is_stale INT,
                          PRIMARY KEY (date, brand_code, product_code));
CREATE TABLE fact_api_snapshot (snapshot_date TEXT, brand_code TEXT, product_code TEXT, price REAL,
                          effective_date TEXT, effective_time TEXT, fetched_at TEXT,
                          PRIMARY KEY (snapshot_date, brand_code, product_code));
CREATE TABLE fact_archive_retail (announce_date TEXT, brand_code TEXT, product_code TEXT, price REAL,
                          effective_date TEXT, effective_time TEXT, source_file TEXT,
                          PRIMARY KEY (announce_date, brand_code, product_code));
CREATE TABLE fact_price_structure (date TEXT, product_code TEXT, product_label TEXT,
                          ex_refinery REAL, refinery_discount REAL, excise_tax REAL, municipal_tax REAL,
                          oil_fund REAL, conservation_fund REAL, wholesale REAL, vat_wholesale REAL,
                          wholesale_incl_vat REAL, marketing_margin REAL, vat_marketing_margin REAL,
                          retail REAL, fx_thb_usd REAL, source_file TEXT,
                          PRIMARY KEY (date, product_code));

-- ready-made views (see sql/queries.sql for more examples)
CREATE VIEW v_monthly AS
  SELECT month, brand_code, product_code, ROUND(AVG(price),3) AS avg_price, MIN(price) AS min_price,
         MAX(price) AS max_price, COUNT(*) AS days
  FROM fact_retail_daily GROUP BY month, brand_code, product_code;
CREATE VIEW v_yearly AS
  SELECT year, brand_code, product_code, ROUND(AVG(price),3) AS avg_price, MIN(price) AS min_price,
         MAX(price) AS max_price, COUNT(*) AS days
  FROM fact_retail_daily GROUP BY year, brand_code, product_code;
CREATE VIEW v_latest AS
  SELECT f.* FROM fact_retail_daily f
  WHERE f.date = (SELECT MAX(date) FROM fact_retail_daily);
CREATE VIEW v_changes AS
  SELECT date, brand_code, product_code, prev_price, price, ROUND(price - prev_price, 2) AS change
  FROM (SELECT date, brand_code, product_code, price,
               LAG(price) OVER (PARTITION BY brand_code, product_code ORDER BY date) AS prev_price,
               LAG(date) OVER (PARTITION BY brand_code, product_code ORDER BY date) AS prev_date
        FROM fact_retail_daily)
  WHERE prev_price IS NOT NULL AND price <> prev_price
    AND julianday(date) - julianday(prev_date) = 1;   -- ignore jumps across data gaps
CREATE VIEW v_vs_ptt AS
  SELECT f.date, f.brand_code, f.product_code, f.price, p.price AS ptt_price,
         ROUND(f.price - p.price, 2) AS diff_vs_ptt
  FROM fact_retail_daily f JOIN fact_retail_daily p
    ON p.date = f.date AND p.product_code = f.product_code AND p.brand_code = 'ptt';
"""


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def build_sqlite(daily, brand_rows, structure_rows, stale, archive_rows=()) -> sqlite3.Connection:
    config.SQLITE_DB.parent.mkdir(parents=True, exist_ok=True)
    if config.SQLITE_DB.exists():
        config.SQLITE_DB.unlink()
    con = sqlite3.connect(config.SQLITE_DB)
    con.executescript(SCHEMA)
    con.executemany("INSERT INTO dim_brand VALUES (?,?,?,?,?,?)", [
        (c, v["th"], v["en"], v["order"], 1 if c in stale else 0, stale.get(c))
        for c, v in config.BRANDS.items()])
    con.executemany("INSERT INTO dim_product VALUES (?,?,?,?,?)", [
        (c, v["th"], v["en"], v["group"], v["order"]) for c, v in config.PRODUCTS.items()])
    con.executemany("INSERT INTO fact_retail_daily VALUES (?,?,?,?,?,?,?,?)", [
        (r["date"], int(r["date"][:4]), r["date"][:7], r["brand_code"], r["product_code"], r["price"],
         r["source"], r["is_stale"]) for r in daily])
    con.executemany("INSERT OR REPLACE INTO fact_api_snapshot VALUES (?,?,?,?,?,?,?)", [
        (r["snapshot_date"], r["brand_code"], r["product_code"], _num(r["price"]), r["effective_date"],
         r["effective_time"], r["fetched_at"]) for r in brand_rows])
    con.executemany("INSERT OR REPLACE INTO fact_archive_retail VALUES (?,?,?,?,?,?,?)", [
        (r["announce_date"], r["brand_code"], r["product_code"], _num(r["price"]), r["effective_date"],
         r["effective_time"], r["source_file"]) for r in archive_rows])
    cols = config.STRUCTURE_COLUMNS
    con.executemany(f"INSERT OR REPLACE INTO fact_price_structure VALUES ({','.join('?' * len(cols))})", [
        tuple(r.get(c) if c in ("date", "product_code", "product_label", "source_file") else _num(r.get(c))
              for c in cols) for r in structure_rows])
    con.commit()
    return con


# ------------------------------------------------------------------ 3/4. pivots
def _brand_cols():
    return sorted(config.BRANDS, key=lambda b: config.BRANDS[b]["order"])


def pivot(con, sql: str, row_keys: list[str], value: str) -> tuple[list[str], list[list]]:
    """Generic pivot: rows = row_keys, columns = brands, cell = value."""
    brands = _brand_cols()
    grid: dict[tuple, dict] = defaultdict(dict)
    for rec in con.execute(sql).fetchall():
        rec = dict(rec)
        grid[tuple(rec[k] for k in row_keys)][rec["brand_code"]] = rec[value]
    header = row_keys + [config.BRANDS[b]["th"] for b in brands]
    rows = [list(k) + [grid[k].get(b) for b in brands] for k in sorted(grid)]
    return header, rows


def _pname(code):
    return config.PRODUCTS.get(code, {}).get("th", code)


def build_pivots(con) -> dict[str, tuple[list[str], list[list]]]:
    con.row_factory = sqlite3.Row
    order = "CASE product_code " + " ".join(
        f"WHEN '{c}' THEN {v['order']}" for c, v in config.PRODUCTS.items()) + " ELSE 99 END"
    out = {}
    h, rows = pivot(con, f"SELECT *, {order} AS po FROM v_latest", ["po", "product_code"], "price")
    out["latest_prices"] = (["product"] + h[2:], [[_pname(r[1])] + r[2:] for r in rows])
    h, rows = pivot(con, f"SELECT *, {order} AS po FROM v_monthly", ["month", "po", "product_code"], "avg_price")
    out["monthly_avg"] = (["month", "product"] + h[3:], [[r[0], _pname(r[2])] + r[3:] for r in rows])
    h, rows = pivot(con, f"SELECT *, {order} AS po FROM v_yearly", ["year", "po", "product_code"], "avg_price")
    out["yearly_avg"] = (["year", "product"] + h[3:], [[r[0], _pname(r[2])] + r[3:] for r in rows])
    h, rows = pivot(con, f"SELECT *, {order} AS po FROM v_monthly", ["month", "po", "product_code"], "max_price")
    out["monthly_max"] = (["month", "product"] + h[3:], [[r[0], _pname(r[2])] + r[3:] for r in rows])
    h, rows = pivot(con, f"SELECT *, {order} AS po FROM v_monthly", ["month", "po", "product_code"], "min_price")
    out["monthly_min"] = (["month", "product"] + h[3:], [[r[0], _pname(r[2])] + r[3:] for r in rows])
    for p in ("gh95", "gh91", "e20", "ds", "dsb20", "gl95", "gs95p", "pds"):
        h, rows = pivot(con, f"SELECT * FROM fact_retail_daily WHERE product_code='{p}'", ["date"], "price")
        out[f"daily_{p}"] = (h, rows)
    rows = con.execute(f"""SELECT c.date, b.name_th AS brand, c.product_code, c.prev_price, c.price, c.change
        FROM v_changes c JOIN dim_brand b USING (brand_code) ORDER BY c.date DESC, b.sort_order, {order}""").fetchall()
    out["price_changes"] = (["date", "brand", "product", "prev_price", "price", "change"],
                            [[r[0], r[1], _pname(r[2]), r[3], r[4], r[5]] for r in rows])
    rows = con.execute(f"""SELECT date, product_code, ex_refinery, excise_tax + municipal_tax AS taxes,
        oil_fund, conservation_fund, marketing_margin, vat_wholesale + vat_marketing_margin AS vat, retail
        FROM fact_price_structure WHERE retail IS NOT NULL ORDER BY date DESC, {order}""").fetchall()
    out["structure_ptt"] = (["date", "product", "ex_refinery", "excise+municipal_tax", "oil_fund",
                             "conservation_fund", "marketing_margin", "vat", "retail"],
                            [[r[0], _pname(r[1])] + list(r[2:]) for r in rows])
    return out


def _safe_cell(v):
    """Prevent CSV/Excel formula injection for text coming from outside."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def write_pivot_csvs(pivots) -> None:
    config.PIVOT_DIR.mkdir(parents=True, exist_ok=True)
    for name, (header, rows) in pivots.items():
        if name.startswith("daily_") or name == "price_changes":
            continue  # large tables live in the Excel/SQLite downloads
        write_csv(config.PIVOT_DIR / f"{name}.csv",
                  [dict(zip(header, [_safe_cell(v) for v in r])) for r in rows], header)


def write_excel(pivots, daily, path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    notes = [
        ["รายงานราคาขายปลีกน้ำมัน (กรุงเทพฯ และปริมณฑล) - ที่มา: สนพ. (EPPO)"],
        ["หน่วย: บาท/ลิตร  |  ค่าเฉลี่ย = เฉลี่ยทุกวันปฏิทินที่มีราคา (ถ่วงตามจำนวนวันที่ราคานั้นมีผล)"],
        ["1 ม.ค. - 10 ก.ค. 2561: ทุกแบรนด์ จากคลังประกาศราคารายแบรนด์ของ สนพ. (ข้อมูลเดียวกับปุ่ม Generate; สนพ. หยุดอัปเดต 10 ก.ค. 2561)"],
        ["11 ก.ค. 2561 - วันก่อนเริ่มเก็บรายวัน: สนพ. มีเฉพาะราคา ปตท. (ไฟล์โครงสร้างราคาน้ำมันรายวัน) -> แบรนด์อื่นเว้นว่าง ไม่มีการเดาตัวเลข"],
        ["ข้อควรระวัง: ค่าเฉลี่ยรายปี 2561 ของแบรนด์อื่นคิดจาก ม.ค.-ก.ค. เท่านั้น (ดูจำนวนวันใน v_yearly.days ของไฟล์ SQLite)"],
        ["source ในชีท Data: eppo_api = ดึงจาก API รายวัน, eppo_api_effective = วันที่ราคานั้นมีผลตาม API,"],
        ["  eppo_structure/seed_structure = จากไฟล์โครงสร้างราคา, eppo_archive = คลังปุ่ม Generate ปี 2561,"],
        ["  carried_forward = ใช้ราคาล่าสุดต่อ (วันหยุด / ราคาจากคลังมีผลจนถึงประกาศครั้งถัดไป)"],
        ["ต้องการ PivotTable เอง: ไปที่ชีท Data > Insert > PivotTable (ข้อมูลเป็น Excel Table ชื่อ tblDaily)"],
    ]
    for n in notes:
        ws.append(n)
    ws.column_dimensions["A"].width = 110

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E79")
    sheet_names = {
        "latest_prices": "Latest", "monthly_avg": "Monthly_Avg", "monthly_min": "Monthly_Min",
        "monthly_max": "Monthly_Max", "yearly_avg": "Yearly_Avg", "price_changes": "Price_Changes",
        "structure_ptt": "Structure_PTT",
    }
    for key, (header, rows) in pivots.items():
        title = sheet_names.get(key) or ("Daily_" + key.split("_", 1)[1])
        s = wb.create_sheet(title[:31])
        s.append(header)
        for r in rows:
            s.append([_safe_cell(v) for v in r])
        for c in s[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        s.freeze_panes = "C2" if header[:2] in (["month", "product"], ["year", "product"]) else "B2"
        s.auto_filter.ref = s.dimensions
        for j in range(1, len(header) + 1):
            s.column_dimensions[get_column_letter(j)].width = 24 if header[j - 1] in ("product", "brand") else (14 if j > 1 else 22)
        for row in s.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "0.00"

    s = wb.create_sheet("Data")
    cols = ["date", "year", "month", "brand_code", "brand", "product_code", "product", "price", "source", "is_stale"]
    s.append(cols)
    for r in daily:
        s.append([r["date"], int(r["date"][:4]), r["date"][:7], r["brand_code"],
                  config.BRANDS.get(r["brand_code"], {}).get("th", r["brand_code"]), r["product_code"],
                  _pname(r["product_code"]), r["price"], r["source"], r["is_stale"]])
    tab = Table(displayName="tblDaily", ref=f"A1:{get_column_letter(len(cols))}{max(2, len(daily) + 1)}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    s.add_table(tab)
    s.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------ 5. dashboard JSON
def build_dashboard_json(con, daily, structure_rows, brand_rows, stale, warnings) -> dict:
    con.row_factory = sqlite3.Row
    dates = sorted({r["date"] for r in daily})
    idx = {d: i for i, d in enumerate(dates)}
    series: dict[str, dict] = {}
    for r in daily:
        k = f"{r['brand_code']}|{r['product_code']}"
        s = series.setdefault(k, {"start": idx[r["date"]], "v": [], "src": ""})
        pos = idx[r["date"]] - s["start"]
        while len(s["v"]) < pos:          # gap (shouldn't happen often) -> null
            s["v"].append(None)
            s["src"] += "-"
        s["v"].append(r["price"])
        s["src"] += SRC_CODE.get(r["source"], "?")

    # monthly / yearly averages are computed in the browser from `series` (same rule as the
    # SQL views v_monthly / v_yearly: mean of daily prices) so they follow the date filter.
    changes = [dict(r) for r in con.execute(
        "SELECT * FROM v_changes ORDER BY date DESC LIMIT 1500")]

    # latest structure (PTT) breakdown
    s_dates = sorted({r["date"] for r in structure_rows if r.get("wholesale") not in ("", None)})
    structure_latest = []
    if s_dates:
        for r in structure_rows:
            if r["date"] == s_dates[-1] and r["product_code"] in config.PRODUCTS:
                structure_latest.append({k: (_num(r.get(k)) if k not in ("date", "product_code") else r[k])
                                         for k in ["date", "product_code", *config.STRUCTURE_FIELDS]})
    # marketing margin & oil fund history (structure), main products
    mm = defaultdict(list)
    for r in structure_rows:
        if r["product_code"] in ("gh95", "gh91", "e20", "ds", "dsb20", "gl95") and r.get("marketing_margin") not in ("", None):
            mm[r["product_code"]].append([r["date"], _num(r["marketing_margin"]), _num(r.get("oil_fund"))])

    first_api = min((r["snapshot_date"] for r in brand_rows), default=None)
    coverage = {}
    for r in daily:
        if r["source"] == "carried_forward":
            continue
        c = coverage.setdefault(r["brand_code"], {"first_real": r["date"], "dates": set(), "sources": set()})
        c["first_real"] = min(c["first_real"], r["date"])
        c["dates"].add(r["date"])
        c["sources"].add(r["source"])
    all_by_brand: dict[str, set] = defaultdict(set)
    for r in daily:
        all_by_brand[r["brand_code"]].add(r["date"])
    for b, c in coverage.items():
        src = c.pop("sources")
        c["real_days"] = len(c.pop("dates"))       # days with an observation (not carried forward)
        c["days"] = len(all_by_brand[b])           # days that have a price
        # contiguous periods that have a price (incl. carried-forward days) -> [[from, to], ...]
        ranges, prev = [], None
        for d in sorted(all_by_brand[b]):
            if prev and (_d(d) - _d(prev)).days == 1:
                ranges[-1][1] = d
            else:
                ranges.append([d, d])
            prev = d
        c["ranges"] = ranges
        parts = []
        if "eppo_archive" in src:
            parts.append("คลังราคารายแบรนด์ของ สนพ. (ปุ่ม Generate) ถึง 10 ก.ค. 2561")
        if src & {"eppo_structure", "seed_structure"}:
            parts.append("ไฟล์โครงสร้างราคาน้ำมันรายวันของ สนพ.")
        if "manual" in src:
            parts.append("ประวัติที่นำเข้าเอง (data/manual)")
        if src & {"eppo_api", "eppo_api_effective"}:
            parts.append("ราคารายวันจาก สนพ.")
        c["history_source"] = " + ".join(parts)
    brand_meta = []
    for code, v in sorted(config.BRANDS.items(), key=lambda kv: kv[1]["order"]):
        keys = [k for k in series if k.startswith(code + "|")]
        brand_meta.append({"code": code, "th": v["th"], "en": v["en"], "color": v["color"],
                           "stale": code in stale, "stale_since": stale.get(code),
                           "has_data": bool(keys)})
    return {
        "meta": {
            "generated_at": config.now_bkk().isoformat(timespec="minutes"),
            "first_date": dates[0] if dates else None,
            "latest_date": dates[-1] if dates else None,
            "per_brand_since": first_api,
            "history_start": config.HISTORY_START,
            "reference_brand": config.STRUCTURE_REFERENCE_BRAND,
            "unit": "บาท/ลิตร",
            "source": "สำนักงานนโยบายและแผนพลังงาน (สนพ./EPPO) - www.eppo.go.th",
            "warnings": warnings[-20:],
        },
        "brands": brand_meta,
        "coverage": coverage,
        "products": [{"code": c, "th": v["th"], "en": v["en"], "group": v["group"]}
                     for c, v in sorted(config.PRODUCTS.items(), key=lambda kv: kv[1]["order"])],
        "dates": dates,
        "series": series,
        "changes": changes,
        "structure_latest": structure_latest,
        "structure_history": mm,
    }


# ------------------------------------------------------------------ 6. daily report
MAIN_PRODUCTS = ["gh95", "gh91", "e20", "gl95", "ds", "dsb20", "pds", "gs95p"]


def build_report_md(con, stale, warnings) -> tuple[str, str]:
    con.row_factory = sqlite3.Row
    latest = con.execute("SELECT MAX(date) FROM fact_retail_daily").fetchone()[0]
    if not latest:
        return "", ""
    prev = (_d(latest) - timedelta(days=1)).isoformat()
    cur = {(r["brand_code"], r["product_code"]): r for r in
           con.execute("SELECT * FROM fact_retail_daily WHERE date=?", (latest,))}
    old = {(r["brand_code"], r["product_code"]): r["price"] for r in
           con.execute("SELECT * FROM fact_retail_daily WHERE date=?", (prev,))}
    lines = [f"# ราคาขายปลีกน้ำมัน ณ วันที่ {latest}", "",
             "หน่วย: บาท/ลิตร · กรุงเทพฯ และปริมณฑล · ที่มา: สนพ. (EPPO)", "",
             "| แบรนด์ | " + " | ".join(_pname(p) for p in MAIN_PRODUCTS) + " |",
             "|---|" + "---:|" * len(MAIN_PRODUCTS)]
    for b in _brand_cols():
        cells = []
        for p in MAIN_PRODUCTS:
            r = cur.get((b, p))
            if not r:
                cells.append("-")
                continue
            ch = ""
            if (b, p) in old and old[(b, p)] != r["price"]:
                d = r["price"] - old[(b, p)]
                ch = f" ({'▲' if d > 0 else '▼'}{abs(d):.2f})"
            cells.append(f"{r['price']:.2f}{ch}")
        name = config.BRANDS[b]["th"] + (" ⚠️ข้อมูลเก่า" if b in stale else "")
        lines.append(f"| {name} | " + " | ".join(cells) + " |")
    changes = [r for r in con.execute("SELECT * FROM v_changes WHERE date=?", (latest,))]
    lines += ["", f"**รายการเปลี่ยนแปลงราคาวันนี้:** {len(changes)} รายการ"]
    for r in changes[:40]:
        lines.append(f"- {config.BRANDS.get(r['brand_code'], {}).get('th', r['brand_code'])} "
                     f"{_pname(r['product_code'])}: {r['prev_price']:.2f} → {r['price']:.2f} ({r['change']:+.2f})")
    if warnings:
        lines += ["", "**ข้อสังเกตคุณภาพข้อมูล (warnings):**"] + [f"- {w}" for w in warnings[:20]]
    return latest, "\n".join(lines) + "\n"


# ------------------------------------------------------------------ orchestration
def run(warnings: list[str] | None = None, end_date: str | None = None) -> dict:
    warnings = list(warnings or [])
    brand_rows = read_csv(config.BRAND_CSV)
    structure_rows = read_csv(config.STRUCTURE_CSV)
    if not brand_rows and not structure_rows:
        raise SystemExit("No data yet - run scripts/fetch_daily.py (or import_seed.py) first.")
    all_dates = [r["snapshot_date"] for r in brand_rows] + [r["date"] for r in structure_rows]
    end_date = end_date or max(max(all_dates), config.today_bkk()) if all_dates else config.today_bkk()
    # never extend past the last real observation by more than the ffill window
    end_date = min(end_date, (_d(max(all_dates)) + timedelta(days=MAX_FFILL_DAYS)).isoformat())

    stale = stale_brands(brand_rows)
    for b, since in stale.items():
        warnings.append(f"{config.BRANDS.get(b, {}).get('th', b)}: EPPO ยังแสดงราคาเก่า (มีผลตั้งแต่ {since}) - "
                        "แสดงไว้แต่ไม่นำไปเทียบ")
    manual_rows, mwarn = read_manual_history()
    warnings += mwarn
    archive_rows = read_csv(config.ARCHIVE_CSV)
    daily = build_daily(brand_rows, structure_rows, end_date, manual_rows, archive_rows)

    # (re)create _site from the committed site/ source
    if config.SITE_OUT.exists():
        shutil.rmtree(config.SITE_OUT)
    shutil.copytree(config.SITE_SRC, config.SITE_OUT)
    (config.SITE_OUT / "data").mkdir(exist_ok=True)
    (config.SITE_OUT / ".nojekyll").write_text("")

    con = build_sqlite(daily, brand_rows, structure_rows, stale, archive_rows)
    pivots = build_pivots(con)
    write_pivot_csvs(pivots)
    write_excel(pivots, daily, config.SITE_OUT / "downloads" / "oil_price_report.xlsx")
    for name in ("latest_prices", "monthly_avg", "yearly_avg"):
        shutil.copy(config.PIVOT_DIR / f"{name}.csv", config.SITE_OUT / "downloads" / f"{name}.csv")
    dash = build_dashboard_json(con, daily, structure_rows, brand_rows, stale, warnings)
    (config.SITE_OUT / "data" / "dashboard.json").write_text(
        json.dumps(dash, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    latest, md = build_report_md(con, stale, warnings)
    if md:
        config.DAILY_REPORT_DIR.mkdir(parents=True, exist_ok=True)
        (config.REPORTS_DIR / "latest.md").write_text(md, encoding="utf-8")
        (config.DAILY_REPORT_DIR / f"{latest}.md").write_text(md, encoding="utf-8")
        summary = os.environ.get("GITHUB_STEP_SUMMARY")
        if summary:
            with open(summary, "a", encoding="utf-8") as f:
                f.write(md)
    con.close()
    log.info("built: %d daily rows, %d dates, latest=%s", len(daily), len(dash["dates"]), latest)
    return {"daily_rows": len(daily), "latest": latest, "stale": stale}
